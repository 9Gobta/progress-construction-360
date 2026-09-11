import uuid
from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import Workbook

from progress_api.services.schedule_import import OPTIONAL_COLUMNS, REQUIRED_COLUMNS


def _schedule_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Task_Table1"
    worksheet.append((*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS))
    worksheet.append(
        (
            "1.งานคานชั้น 1",
            "1",
            "January 15, 2026 8:00 AM",
            "January 31, 2026 5:00 PM",
            "NA",
            "NA",
            "1",
        )
    )
    worksheet.append(
        (
            "1.1.งานผูกเหล็กคานชั้น 1",
            "1.1",
            "January 17, 2026 8:00 AM",
            "January 23, 2026 5:00 PM",
            "NA",
            "NA",
            "0.5",
        )
    )
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _revised_schedule_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Task_Table1"
    worksheet.append((*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS))
    worksheet.append((
        "1.งานคานชั้น 1 (แผนใหม่)", "1",
        "January 20, 2026 8:00 AM", "February 5, 2026 5:00 PM", "NA", "NA", "1",
    ))
    worksheet.append((
        "1.1.งานผูกเหล็กคานชั้น 1 (แผนใหม่)", "1.1",
        "January 22, 2026 8:00 AM", "January 30, 2026 5:00 PM", "NA", "NA", "0.5",
    ))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _ground_beam_schedule_bytes() -> bytes:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Task_Table1"
    worksheet.append((*REQUIRED_COLUMNS, *OPTIONAL_COLUMNS))
    worksheet.append((
        "งานคานคอดินชั้น 1", "1.2.1.2",
        "December 20, 2025 8:00 AM", "January 31, 2026 5:00 PM",
        "NA", "NA", "1",
    ))
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _project(client: TestClient, headers: dict[str, str]) -> str:
    response = client.post(
        "/api/v1/projects",
        headers=headers,
        json={"name": "Schedule Test", "timezone": "Asia/Bangkok"},
    )
    assert response.status_code == 201
    return response.json()["id"]


def test_progress_ai_endpoints_are_retired(
    client: TestClient,
    auth_headers: dict[str, str],
) -> None:
    project_id = _project(client, auth_headers)
    capture_id = uuid.uuid4()
    floor_id = uuid.uuid4()

    evaluation = client.get(
        f"/api/v1/projects/{project_id}/beam-ai/evaluation",
        headers=auth_headers,
    )
    beam_run = client.post(
        f"/api/v1/projects/{project_id}/captures/{capture_id}/beam-ai/run",
        headers=auth_headers,
        json={"floor_id": str(floor_id)},
    )
    work_run = client.post(
        f"/api/v1/projects/{project_id}/captures/{capture_id}/work-ai/run",
        headers=auth_headers,
        json={"floor_id": str(floor_id)},
    )

    assert evaluation.status_code == 410
    assert beam_run.status_code == 410
    assert work_run.status_code == 410
    assert "ยกเลิก" in evaluation.json()["detail"]


def test_manual_progress_bulk_saves_every_selected_activity_atomically(
    client: TestClient,
    auth_headers: dict[str, str],
    monkeypatch,
) -> None:
    project_id = _project(client, auth_headers)
    monkeypatch.setattr("progress_api.api.routes.schedules.put_object", lambda **_: None)
    imported = client.post(
        f"/api/v1/projects/{project_id}/schedules/import",
        headers=auth_headers,
        data={"name": "Bulk progress", "is_baseline": "true"},
        files={"file": ("schedule.xlsx", _schedule_bytes(), "application/octet-stream")},
    )
    assert imported.status_code == 201
    activities = client.get(
        f"/api/v1/projects/{project_id}/schedules/{imported.json()['id']}/activities",
        headers=auth_headers,
    ).json()

    saved = client.post(
        f"/api/v1/projects/{project_id}/progress/manual/bulk",
        headers=auth_headers,
        json={"entries": [{
            "activity_id": activity["id"],
            "observed_at": "2026-01-20T10:00:00+07:00",
            "progress_percent": 75,
            "note": "บันทึกพร้อมกันจากภาพ 360",
        } for activity in activities]},
    )
    assert saved.status_code == 201
    assert len(saved.json()) == len(activities)
    assert {row["activity_wbs"] for row in saved.json()} == {
        activity["wbs"] for activity in activities
    }
    assert {row["progress_percent"] for row in saved.json()} == {"75.000"}

    duplicate = client.post(
        f"/api/v1/projects/{project_id}/progress/manual/bulk",
        headers=auth_headers,
        json={"entries": [{
            "activity_id": activities[0]["id"],
            "observed_at": "2026-01-20T10:00:00+07:00",
            "progress_percent": value,
        } for value in (25, 50)]},
    )
    assert duplicate.status_code == 422
    history = client.get(
        f"/api/v1/projects/{project_id}/progress/manual",
        headers=auth_headers,
    )
    assert len(history.json()) == len(activities)


def test_schedule_preview_import_and_manual_actual(
    client: TestClient,
    auth_headers: dict[str, str],
    monkeypatch,
) -> None:
    project_id = _project(client, auth_headers)
    content = _schedule_bytes()
    uploaded: list[str] = []
    monkeypatch.setattr(
        "progress_api.api.routes.schedules.put_object",
        lambda *, key, body, content_type: uploaded.append(key),
    )

    preview = client.post(
        f"/api/v1/projects/{project_id}/schedules/preview",
        headers=auth_headers,
        files={"file": ("schedule.xlsx", content, "application/octet-stream")},
    )
    assert preview.status_code == 200
    assert preview.json()["row_count"] == 2
    assert len(preview.json()["warnings"]) == 2
    assert "Percent_Complete was ignored" in preview.json()["warnings"][1]

    imported = client.post(
        f"/api/v1/projects/{project_id}/schedules/import",
        headers=auth_headers,
        data={"name": "Baseline v1", "is_baseline": "true"},
        files={"file": ("schedule.xlsx", content, "application/octet-stream")},
    )
    assert imported.status_code == 201
    schedule = imported.json()
    assert schedule["is_baseline"] is True
    assert schedule["row_count"] == 2
    assert len(uploaded) == 1

    duplicate = client.post(
        f"/api/v1/projects/{project_id}/schedules/import",
        headers=auth_headers,
        data={"name": "Duplicate", "is_baseline": "true"},
        files={"file": ("schedule.xlsx", content, "application/octet-stream")},
    )
    assert duplicate.status_code == 409

    activities = client.get(
        f"/api/v1/projects/{project_id}/schedules/{schedule['id']}/activities",
        headers=auth_headers,
    )
    assert activities.status_code == 200
    activity_rows = activities.json()
    assert len(activity_rows) == 2
    assert "percent_complete_source" not in activity_rows[0]

    actual = client.post(
        f"/api/v1/projects/{project_id}/progress/manual",
        headers=auth_headers,
        json={
            "activity_id": activity_rows[1]["id"],
            "observed_at": "2025-12-28T17:00:00+07:00",
            "progress_percent": 40,
            "note": "ตรวจจากหน้างาน",
        },
    )
    assert actual.status_code == 201
    assert actual.json()["progress_percent"] == "40.000"

    newer_actual = client.post(
        f"/api/v1/projects/{project_id}/progress/manual",
        headers=auth_headers,
        json={
            "activity_id": activity_rows[1]["id"],
            "observed_at": "2025-12-28T17:00:00+07:00",
            "progress_percent": 50,
            "note": "แก้หลังตรวจหลักฐานอีกครั้ง",
        },
    )
    assert newer_actual.status_code == 201

    history = client.get(
        f"/api/v1/projects/{project_id}/progress/manual",
        headers=auth_headers,
    )
    assert history.status_code == 200
    assert len(history.json()) == 2
    assert {row["progress_percent"] for row in history.json()} == {"40.000", "50.000"}
    assert {row["activity_wbs"] for row in history.json()} == {"1.1"}

    floor = client.post(
        f"/api/v1/projects/{project_id}/floors",
        headers=auth_headers,
        json={"name": "ชั้น 1", "level_index": 1},
    ).json()
    media = client.post(
        f"/api/v1/projects/{project_id}/media",
        headers=auth_headers,
        json={
            "original_filename": "comparison.mp4",
            "content_type": "video/mp4",
            "size_bytes": 1024,
        },
    ).json()
    capture = client.post(
        f"/api/v1/projects/{project_id}/captures",
        headers=auth_headers,
        json={
            "source_video_id": media["id"],
            "captured_at": "2025-12-28T18:00:00+07:00",
            "start_floor_id": floor["id"],
            "start_x": 0.5,
            "start_y": 0.5,
        },
    ).json()
    comparison = client.get(
        f"/api/v1/projects/{project_id}/progress/comparison",
        headers=auth_headers,
        params={"capture_id": capture["id"]},
    )
    assert comparison.status_code == 200
    item = next(row for row in comparison.json()["items"] if row["wbs"] == "1.1")
    assert item["planned_percent"] == "0"
    assert item["human_actual_percent"] == "50.000"
    assert item["variance_pp"] == "50.000"

    revised = client.post(
        f"/api/v1/projects/{project_id}/schedules/import",
        headers=auth_headers,
        data={"name": "Baseline v2", "is_baseline": "true"},
        files={
            "file": (
                "schedule-v2.xlsx",
                _revised_schedule_bytes(),
                "application/octet-stream",
            )
        },
    )
    assert revised.status_code == 201
    revised_activities = client.get(
        f"/api/v1/projects/{project_id}/schedules/{revised.json()['id']}/activities",
        headers=auth_headers,
    ).json()
    revised_detail = next(row for row in revised_activities if row["wbs"] == "1.1")
    assert revised_detail["id"] != activity_rows[1]["id"]

    revised_comparison = client.get(
        f"/api/v1/projects/{project_id}/progress/comparison",
        headers=auth_headers,
        params={"capture_id": capture["id"]},
    )
    assert revised_comparison.status_code == 200
    revised_item = next(
        row for row in revised_comparison.json()["items"] if row["wbs"] == "1.1"
    )
    assert revised_item["activity_id"] == revised_detail["id"]
    assert revised_item["human_actual_percent"] == "50.000"


def test_comparison_carries_beam_progress_forward_and_counts_uninspected_beams(
    client: TestClient,
    auth_headers: dict[str, str],
    monkeypatch,
) -> None:
    project_id = _project(client, auth_headers)
    monkeypatch.setattr(
        "progress_api.api.routes.schedules.put_object",
        lambda **_: None,
    )
    imported = client.post(
        f"/api/v1/projects/{project_id}/schedules/import",
        headers=auth_headers,
        data={"name": "Ground beam baseline", "is_baseline": "true"},
        files={
            "file": (
                "ground-beam.xlsx",
                _ground_beam_schedule_bytes(),
                "application/octet-stream",
            )
        },
    )
    assert imported.status_code == 201

    floor = client.post(
        f"/api/v1/projects/{project_id}/floors",
        headers=auth_headers,
        json={"name": "ชั้น 1", "level_index": 1},
    ).json()
    earlier_media = client.post(
        f"/api/v1/projects/{project_id}/media",
        headers=auth_headers,
        json={
            "original_filename": "beam-comparison.mp4",
            "content_type": "video/mp4",
            "size_bytes": 1024,
        },
    ).json()
    later_media = client.post(
        f"/api/v1/projects/{project_id}/media",
        headers=auth_headers,
        json={
            "original_filename": "beam-comparison-next-day.mp4",
            "content_type": "video/mp4",
            "size_bytes": 1024,
        },
    ).json()
    earlier = client.post(
        f"/api/v1/projects/{project_id}/captures",
        headers=auth_headers,
        json={
            "source_video_id": earlier_media["id"],
            "captured_at": "2025-12-24T10:00:00+07:00",
            "start_floor_id": floor["id"],
            "start_x": 0.2,
            "start_y": 0.4,
        },
    ).json()
    later = client.post(
        f"/api/v1/projects/{project_id}/captures",
        headers=auth_headers,
        json={
            "source_video_id": later_media["id"],
            "captured_at": "2025-12-25T10:00:00+07:00",
            "start_floor_id": floor["id"],
            "start_x": 0.2,
            "start_y": 0.4,
        },
    ).json()
    segments = client.put(
        f"/api/v1/projects/{project_id}/floors/{floor['id']}/beam-segments",
        headers=auth_headers,
        json={
            "sheet_name": "ST-03",
            "segments": [
                {
                    "code": "B-1",
                    "beam_type": "GB",
                    "start_x": 0.2,
                    "start_y": 0.4,
                    "end_x": 0.3,
                    "end_y": 0.4,
                    "length_m": 2,
                },
                {
                    "code": "B-2",
                    "beam_type": "GB",
                    "start_x": 0.3,
                    "start_y": 0.4,
                    "end_x": 0.4,
                    "end_y": 0.4,
                    "length_m": 2,
                },
            ],
        },
    ).json()
    saved = client.put(
        f"/api/v1/projects/{project_id}/captures/{earlier['id']}/beam-progress",
        headers=auth_headers,
        json={
            "floor_id": floor["id"],
            "entries": [{"beam_segment_id": segments[0]["id"], "progress_percent": 100}],
        },
    )
    assert saved.status_code == 200

    comparison = client.get(
        f"/api/v1/projects/{project_id}/progress/comparison",
        headers=auth_headers,
        params={"capture_id": later["id"]},
    )
    assert comparison.status_code == 200
    item = next(row for row in comparison.json()["items"] if row["wbs"] == "1.2.1.2")
    assert item["human_actual_percent"] == "50.000"
    assert item["human_observed_at"].startswith("2025-12-24T10:00:00")
