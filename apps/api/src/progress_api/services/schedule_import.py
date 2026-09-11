from __future__ import annotations

from dataclasses import dataclass, replace
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, BinaryIO

from openpyxl import load_workbook

REQUIRED_COLUMNS = (
    "Name",
    "WBS",
    "Start_Date",
    "Finish_Date",
    "Baseline_Start",
    "Baseline_Finish",
)
OPTIONAL_COLUMNS = ("Percent_Complete",)
DATETIME_FORMATS = (
    "%B %d, %Y %I:%M %p",
    "%Y-%m-%d %H:%M:%S",
    "%Y-%m-%d",
)
EMPTY_MARKERS = {"", "NA", "N/A", "NONE", "NULL", "-"}
UPPER_FLOOR_SHORING_GROUPS = {"1": "คาน", "2": "พื้น", "4": "บันได"}


class ScheduleImportError(ValueError):
    """Raised when a workbook cannot be mapped without guessing."""


@dataclass(frozen=True)
class ImportedActivity:
    source_row_no: int
    name: str
    wbs: str
    parent_wbs: str | None
    planned_start: datetime
    planned_finish: datetime
    is_summary: bool


@dataclass(frozen=True)
class ScheduleImportResult:
    sheet_name: str
    rows: tuple[ImportedActivity, ...]
    used_baseline_dates: bool
    warnings: tuple[str, ...]


def _is_empty(value: Any) -> bool:
    return value is None or str(value).strip().upper() in EMPTY_MARKERS


def parse_datetime(value: Any, *, column: str, row_no: int) -> datetime:
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for date_format in DATETIME_FORMATS:
        try:
            return datetime.strptime(text, date_format)
        except ValueError:
            continue
    raise ScheduleImportError(f"Row {row_no}: invalid {column} date: {text!r}")


def parent_wbs(wbs: str) -> str | None:
    parent, separator, _ = wbs.rpartition(".")
    return parent if separator else None


def add_upper_floor_shoring_activities(
    rows: tuple[ImportedActivity, ...],
) -> tuple[tuple[ImportedActivity, ...], bool]:
    """Add the reviewed shoring step missing from upper-floor schedule files."""
    result = list(rows)
    changed = False
    for level in (2, 3, 4):
        for group_suffix, group_label in UPPER_FLOOR_SHORING_GROUPS.items():
            group_wbs = f"1.2.{level}.{group_suffix}"
            direct = [
                (index, row)
                for index, row in enumerate(result)
                if row.parent_wbs == group_wbs and not row.is_summary
            ]
            if not direct or any("ค้ำยัน" in row.name for _index, row in direct):
                continue
            shifted_wbs = {
                row.wbs: f"{group_wbs}.{int(row.wbs.rsplit('.', 1)[1]) + 1}"
                for _index, row in direct
                if row.wbs.rsplit(".", 1)[1].isdigit()
            }
            if len(shifted_wbs) != len(direct):
                continue
            result = [
                replace(row, wbs=shifted_wbs.get(row.wbs, row.wbs))
                for row in result
            ]
            first_index, first = direct[0]
            result.insert(first_index, ImportedActivity(
                source_row_no=first.source_row_no,
                name=f"งานค้ำยัน{group_label}ชั้น {level}",
                wbs=f"{group_wbs}.1",
                parent_wbs=group_wbs,
                planned_start=first.planned_start,
                planned_finish=first.planned_finish,
                is_summary=False,
            ))
            changed = True
    return tuple(result), changed


def load_schedule_workbook(path: str | Path | bytes | BinaryIO) -> ScheduleImportResult:
    source: str | Path | BinaryIO
    source = BytesIO(path) if isinstance(path, bytes) else path
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if len(workbook.sheetnames) != 1:
            raise ScheduleImportError("Workbook must contain exactly one task sheet")
        worksheet = workbook[workbook.sheetnames[0]]
        rows = worksheet.iter_rows(values_only=True)
        try:
            headers = tuple(str(value).strip() if value is not None else "" for value in next(rows))
        except StopIteration as exc:
            raise ScheduleImportError("Workbook is empty") from exc
        required_header_count = len(REQUIRED_COLUMNS)
        if headers[:required_header_count] != REQUIRED_COLUMNS or any(
            header not in OPTIONAL_COLUMNS for header in headers[required_header_count:]
        ):
            raise ScheduleImportError(
                f"Expected required columns {REQUIRED_COLUMNS!r}, received {headers!r}"
            )

        raw_rows = list(rows)
        imported: list[ImportedActivity] = []
        used_baseline_for_any_row = False
        warnings: list[str] = []
        wbs_values = [str(row[1]).strip() for row in raw_rows if len(row) >= 2 and row[1]]

        for row_no, row in enumerate(raw_rows, start=2):
            if not any(value is not None and str(value).strip() for value in row):
                continue
            name = str(row[0]).strip()
            wbs = str(row[1]).strip()
            if not name or not wbs:
                raise ScheduleImportError(f"Row {row_no}: Name and WBS are required")

            has_baseline = not _is_empty(row[4]) and not _is_empty(row[5])
            start_value = row[4] if has_baseline else row[2]
            finish_value = row[5] if has_baseline else row[3]
            used_baseline_for_any_row = used_baseline_for_any_row or has_baseline
            planned_start = parse_datetime(start_value, column="planned start", row_no=row_no)
            planned_finish = parse_datetime(finish_value, column="planned finish", row_no=row_no)
            if planned_finish < planned_start:
                raise ScheduleImportError(f"Row {row_no}: Finish is before Start")

            imported.append(
                ImportedActivity(
                    source_row_no=row_no,
                    name=name,
                    wbs=wbs,
                    parent_wbs=parent_wbs(wbs),
                    planned_start=planned_start,
                    planned_finish=planned_finish,
                    is_summary=any(candidate.startswith(f"{wbs}.") for candidate in wbs_values),
                )
            )

        if not used_baseline_for_any_row:
            warnings.append(
                "Baseline_Start/Baseline_Finish are empty; Start_Date/Finish_Date were used."
            )
        if "Percent_Complete" in headers:
            warnings.append(
                "Percent_Complete was ignored; Human Actual must be entered "
                "and audited in the Web app."
            )
        normalized_rows, added_shoring = add_upper_floor_shoring_activities(tuple(imported))
        if added_shoring:
            warnings.append(
                "เพิ่มขั้นงานค้ำยันสำหรับคาน พื้น และบันไดชั้น 2–4 ตามรูปแบบตรวจของโครงการ"
            )
        return ScheduleImportResult(
            sheet_name=worksheet.title,
            rows=normalized_rows,
            used_baseline_dates=used_baseline_for_any_row,
            warnings=tuple(warnings),
        )
    finally:
        workbook.close()
